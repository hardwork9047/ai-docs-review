"""Read slide text and font families from .pptx bytes with python-pptx (in memory).

LibreOffice の PDF では日本語テキストが化けたり書体が置換されたりするため、
pptx のアップロードではテキストと書体名をここで元ファイルから読む。
"""

import io
import zipfile

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide

from app.domain.pages import SlideText
from app.infra.errors import DocumentError


def read_slide_texts(data: bytes) -> list[SlideText]:
    """Return one `SlideText` per *visible* slide (hidden slides are not exported to PDF).

    Fonts are the typefaces set explicitly on runs plus the theme's body (and, when the
    slide has a title, heading) fonts, with theme references like "+mn-ea" resolved.
    Raise `DocumentError` when `data` is not a valid .pptx package.
    """
    try:
        prs = Presentation(io.BytesIO(data))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as exc:
        raise DocumentError(".pptx として読み込めませんでした") from exc
    theme = _theme_fonts(prs.slide_masters[0].part.part_related_by(RT.THEME).blob)
    return [_read_slide(s, theme) for s in prs.slides if s._element.get("show") != "0"]


def _theme_fonts(theme_xml: bytes) -> dict[str, str]:
    """Map theme references (+mj-lt / +mj-ea / +mn-lt / +mn-ea) to typefaces."""
    root = parse_xml(theme_xml)
    fonts: dict[str, str] = {}
    for tag, prefix in (("a:majorFont", "+mj"), ("a:minorFont", "+mn")):
        scheme = root.find(f".//{qn(tag)}")
        if scheme is None:
            continue
        latin = scheme.find(qn("a:latin"))
        ea = scheme.find(qn("a:ea"))
        jpan = next((f for f in scheme.iter(qn("a:font")) if f.get("script") == "Jpan"), None)
        fonts[f"{prefix}-lt"] = latin.get("typeface", "") if latin is not None else ""
        # 日本語の書体は a:ea が空で a:font script="Jpan" に入っていることが多い
        ea_face = ea.get("typeface", "") if ea is not None else ""
        fonts[f"{prefix}-ea"] = ea_face or (jpan.get("typeface", "") if jpan is not None else "")
    return fonts


def _read_slide(slide: Slide, theme: dict[str, str]) -> SlideText:
    title_shape = slide.shapes.title
    title_id = title_shape.shape_id if title_shape is not None else None
    title, bodies = "", []
    for shape in slide.shapes:
        if isinstance(shape, Shape) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape.shape_id == title_id:
                title = text
            else:
                bodies.append(text)
        elif isinstance(shape, GraphicFrame) and shape.has_table:
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
            bodies.append("[表]\n" + "\n".join(rows))

    notes = ""
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""

    faces = {e.get("typeface", "") for e in slide._element.iter(qn("a:latin"), qn("a:ea"))}
    if bodies:
        faces |= {"+mn-lt", "+mn-ea"}
    if title:
        faces |= {"+mj-lt", "+mj-ea"}
    fonts = sorted({theme.get(f, f) if f.startswith("+") else f for f in faces} - {""})
    return SlideText(title=title, body="\n".join(bodies), notes=notes, fonts=tuple(fonts))
