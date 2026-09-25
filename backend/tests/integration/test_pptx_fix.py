"""Regression tests for app.infra.pptx_fix.

Bug: テーマの a:ea が空だと LibreOffice の PDF で日本語が文字化けする(pptx_fix の docstring 参照)。
"""

import io
import re
import zipfile

from pptx import Presentation

from app.infra.pptx_fix import fill_theme_east_asian_fonts


def _pptx() -> bytes:
    # python-pptx の既定テンプレートも a:ea="" + Jpan="ＭＳ Ｐゴシック" の構成
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "タイトル"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _theme(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("ppt/theme/theme1.xml").decode()


def _ea_faces(xml: str) -> list[str]:
    return re.findall(r'<a:ea typeface="([^"]*)"', xml)


def test_default_theme_really_has_empty_east_asian_fonts() -> None:
    assert _ea_faces(_theme(_pptx())) == ["", ""]


def test_empty_east_asian_fonts_are_filled_from_the_jpan_script_font() -> None:
    fixed = fill_theme_east_asian_fonts(_pptx())
    assert _ea_faces(_theme(fixed)) == ["ＭＳ Ｐゴシック", "ＭＳ Ｐゴシック"]


def test_fixed_package_is_still_a_valid_presentation_with_same_entries() -> None:
    original = _pptx()
    fixed = fill_theme_east_asian_fonts(original)
    assert Presentation(io.BytesIO(fixed)).slides[0].shapes.title.text == "タイトル"
    with zipfile.ZipFile(io.BytesIO(original)) as a, zipfile.ZipFile(io.BytesIO(fixed)) as b:
        assert a.namelist() == b.namelist()
        assert a.read("ppt/slides/slide1.xml") == b.read("ppt/slides/slide1.xml")


def test_explicit_east_asian_font_is_left_alone() -> None:
    once = fill_theme_east_asian_fonts(_pptx())
    filled, explicit = '<a:ea typeface="ＭＳ Ｐゴシック"', '<a:ea typeface="游ゴシック"'
    theme = _theme(once).replace(filled, explicit, 1)
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(once)) as src, zipfile.ZipFile(buf, "w") as dst:
        for name in src.namelist():
            body = theme.encode() if name == "ppt/theme/theme1.xml" else src.read(name)
            dst.writestr(name, body)
    assert _ea_faces(_theme(fill_theme_east_asian_fonts(buf.getvalue())))[0] == "游ゴシック"


def test_non_zip_bytes_are_returned_unchanged() -> None:
    assert fill_theme_east_asian_fonts(b"not a zip") == b"not a zip"
