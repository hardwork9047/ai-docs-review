"""Integration tests for app.infra.document — 形式判定と pptx/PDF の読み込み経路。"""

import io

import pytest
from pptx import Presentation

from app.infra.document import detect_kind, load_document
from app.infra.errors import DocumentError
from tests.pdf_factory import make_pdf

PDF = make_pdf([[("Rendered", 24)], [("Second", 24)]])


def _pptx(*titles: str) -> bytes:
    prs = Presentation()
    for title in titles:
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    ("name", "data", "kind"),
    [("a.pdf", b"%PDF-1.7", "pdf"), ("A.PPTX", b"PK\x03\x04", "pptx")],
)
def test_detect_kind_accepts_pdf_and_pptx(name: str, data: bytes, kind: str) -> None:
    assert detect_kind(name, data) == kind


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("a.ppt", b"PK"),
        ("a.docx", b"PK"),
        ("a.png", b"\x89PNG"),
        ("a.pdf", b"PK\x03\x04"),  # 拡張子と中身が食い違う
        ("a.pptx", b"%PDF"),
        ("noext", b"%PDF"),
    ],
)
def test_detect_kind_rejects_everything_else(name: str, data: bytes) -> None:
    with pytest.raises(DocumentError):
        detect_kind(name, data)


def test_pdf_is_read_directly_without_conversion() -> None:
    def must_not_convert(data: bytes, soffice: str) -> bytes:
        raise AssertionError("PDF must not be converted")

    pages = load_document("deck.pdf", PDF, soffice="x", convert=must_not_convert)
    assert [p.title for p in pages] == ["Rendered", "Second"]


def test_pptx_is_converted_and_text_comes_from_the_pptx() -> None:
    calls: list[str] = []

    def fake_convert(data: bytes, soffice: str) -> bytes:
        calls.append(soffice)
        return PDF

    pages = load_document(
        "deck.pptx", _pptx("承認依頼", "結論"), soffice="/opt/soffice", convert=fake_convert
    )
    assert calls == ["/opt/soffice"]
    assert [p.title for p in pages] == ["承認依頼", "結論"]  # pptx 由来
    assert pages[0].char_sizes  # 計測値は PDF 由来
    assert pages[0].image
