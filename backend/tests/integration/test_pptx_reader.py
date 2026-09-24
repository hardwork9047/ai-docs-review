"""Integration tests for app.infra.pptx_reader — python-pptx で生成した pptx を読む。"""

import io

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from app.domain.pages import SlideText
from app.infra.errors import DocumentError
from app.infra.pptx_reader import read_slide_texts

TITLE_AND_CONTENT = 1
BLANK = 6


def _to_bytes(prs: Presentation) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _deck() -> bytes:
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    s1.shapes.title.text = "設備投資の承認依頼"
    body = s1.placeholders[1].text_frame
    body.text = "結論: A案を採用"
    run = body.paragraphs[0].runs[0]
    run.font.name = "Arial"
    ea = run.font._rPr.get_or_add_latin()  # 同じ rPr に東アジア書体も付ける
    ea_el = ea.makeelement(qn("a:ea"), {"typeface": "メイリオ"})
    ea.addnext(ea_el)
    s1.notes_slide.notes_text_frame.text = "冒頭で結論を述べる"

    hidden = prs.slides.add_slide(prs.slide_layouts[BLANK])
    hidden._element.set("show", "0")

    s3 = prs.slides.add_slide(prs.slide_layouts[BLANK])
    table = s3.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    for r, row in enumerate([["項目", "金額"], ["設備", "500万円"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text

    return _to_bytes(prs)


def test_hidden_slides_are_skipped() -> None:
    assert len(read_slide_texts(_deck())) == 2


def test_title_body_and_notes_are_extracted() -> None:
    slide = read_slide_texts(_deck())[0]
    assert (slide.title, slide.body, slide.notes) == (
        "設備投資の承認依頼",
        "結論: A案を採用",
        "冒頭で結論を述べる",
    )


def test_explicit_and_theme_fonts_are_collected() -> None:
    fonts = set(read_slide_texts(_deck())[0].fonts)
    assert {"Arial", "メイリオ", "Calibri"} <= fonts  # Calibri は既定テーマの本文書体
    assert not any(f.startswith("+") for f in fonts)


def test_table_is_rendered_as_pipe_rows() -> None:
    slide = read_slide_texts(_deck())[1]
    assert slide == SlideText(title="", body="[表]\n項目 | 金額\n設備 | 500万円", fonts=slide.fonts)


def test_non_pptx_bytes_raise_document_error() -> None:
    with pytest.raises(DocumentError):
        read_slide_texts(b"not a pptx")
