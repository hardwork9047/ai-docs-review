"""Integration tests: hit the API boundary with a real test client and a fake LLM."""

import io
import json
from collections.abc import Iterator
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api.routes import get_llm
from app.domain.service import LLMError
from app.infra.ollama import OllamaHealth
from app.main import create_app

REPLY = json.dumps({"score": 90, "summary": "s", "good_points": ["g"], "issues": []})
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class FakeLLM:
    """Stands in for OllamaClient: canned chat reply (or error) and canned health."""

    def __init__(self, error: Exception | None = None) -> None:
        self.error = error

    async def complete(self, system: str, user: str, schema: dict[str, Any]) -> str:
        if self.error:
            raise self.error
        return REPLY

    async def health(self) -> OllamaHealth:
        return OllamaHealth(ok=True, model="fake:1", model_ready=True)


@pytest.fixture
def llm() -> FakeLLM:
    return FakeLLM()


@pytest.fixture
def client(llm: FakeLLM) -> Iterator[TestClient]:
    app = create_app()
    app.dependency_overrides[get_llm] = lambda: llm
    with TestClient(app) as c:
        yield c


def _pptx(*titles: str) -> bytes:
    prs = Presentation()
    for title in titles:
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _post(client: TestClient, data: bytes, name: str = "deck.pptx") -> Any:
    return client.post("/api/review", files={"file": (name, data, PPTX_TYPE)})


def _lines(response: Any) -> list[dict[str, Any]]:
    return [json.loads(line) for line in response.text.splitlines()]


def test_health_reports_ollama_status(client: TestClient) -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"ok": True, "model": "fake:1", "model_ready": True, "error": None}


def test_review_streams_meta_then_result_as_ndjson(client: TestClient) -> None:
    response = _post(client, _pptx("承認依頼", "結論"))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/x-ndjson")
    meta, result = _lines(response)
    assert meta["type"] == "meta"
    assert meta["slide_count"] == 2
    assert meta["reviewer"]["name"] == "上司(部長)"
    assert "system_prompt" not in meta["reviewer"]
    assert result["type"] == "result"
    assert result["review"]["score"] == 90
    assert result["verdict"]["overall_passed"] is True


def test_llm_failure_is_streamed_as_error_event(client: TestClient, llm: FakeLLM) -> None:
    llm.error = LLMError("down")
    events = _lines(_post(client, _pptx("t")))
    assert [e["type"] for e in events] == ["meta", "error"]


def test_non_pptx_filename_is_400(client: TestClient) -> None:
    response = _post(client, b"x", name="deck.pdf")
    assert response.status_code == 400


def test_corrupt_pptx_is_400(client: TestClient) -> None:
    response = _post(client, b"not a pptx")
    assert response.status_code == 400


def test_deck_without_slides_is_400(client: TestClient) -> None:
    response = _post(client, _pptx())
    assert response.status_code == 400
