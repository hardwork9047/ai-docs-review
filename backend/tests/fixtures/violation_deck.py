"""Sample proposal deck with 20 planted rule violations, and its answer key.

架空の会社(サンプルケア株式会社)の介護施設向け提案書。`packs/sample_care_sales.yaml` と
組み込みの機械チェックに対して、ちょうど EXPECTED の 20 件が検出される。
Copilot など他のツールとの比較実験や、デモにも使う(`make sample-deck` で pptx を出力)。
"""

import sys
from pathlib import Path

PACK_PATH = Path(__file__).parent / "packs" / "sample_care_sales.yaml"

# (タイトル, 本文) × 9 ページ。違反はコメントの番号で EXPECTED と対応する
SLIDES: list[tuple[str, str]] = [
    (
        "見守りシステムのご提案",
        "サンプル介護ソフト株式会社 営業部\n社外秘",  # 1 旧社名, 2 社外秘
    ),
    (
        "現状の課題",
        # 3 入居者さん, 4 利用者さん
        "夜間の巡回で入居者さんの転倒に気づくのが遅れる\n利用者さんの見守りに人手が足りない",
    ),
    (
        "ご提案の概要",
        "センサーで離床を検知し、スタッフのスマートフォンに通知します\n"
        "夜間の事故は必ず減ります\n業界No.1の検知精度",  # 5 必ず, 6 業界No.1
    ),
    (
        "主な機能",
        "御社の居室にｾﾝｻｰを設置します\n通知はスマートフォンで受け取れます",  # 7 御社, 8 半角カナ
    ),
    (
        "導入の効果",
        "転倒の見逃しを100%防ぎます\n夜勤の負担が確実に軽くなります",  # 9 100%, 10 確実に
    ),
    (
        "導入事例",
        # 11 個人情報, 12 日本一
        "A施設では山田花子様(85歳)の転倒を未然に防ぎました\n日本一の導入実績",
    ),
    (
        "支援体制",
        "専用のサーバーで記録を保管します\nご不明な点はご相談下さい",  # 13/14 表記ゆれ(後半)
    ),
    (
        "運用について",
        "サーバの保守は当社が行います\n社内限:保守の単価表は別途",  # 13 サーバー/サーバ, 15 社内限
    ),
    (
        "まとめ",
        # 16 絶対, 17 最高の, 14 下さい/ください(前半)
        "絶対に後悔させません\n最高の見守りをご提供します\nぜひご検討ください",
    ),
]
# 料金(18)・スケジュール(19)・問い合わせ先(20)は、どのページにも書かない

# (ルール ID または組み込みチェックの種類, ページ番号。0 は資料全体)
EXPECTED: list[tuple[str, int]] = [
    ("R-NAME-01", 1),
    ("R-CONF-01", 1),
    ("R-TERM-01", 2),
    ("R-TERM-01", 2),
    ("R-EXPR-01", 3),
    ("R-EXPR-02", 3),
    ("R-TERM-02", 4),
    ("半角カナ", 0),
    ("R-EXPR-01", 5),
    ("R-EXPR-01", 5),
    ("R-PRIV-01", 6),
    ("R-EXPR-02", 6),
    ("表記ゆれ", 0),
    ("表記ゆれ", 0),
    ("R-CONF-01", 8),
    ("R-EXPR-01", 9),
    ("R-EXPR-02", 9),
    ("R-REQ-01", 0),
    ("R-REQ-02", 0),
    ("R-REQ-03", 0),
]


def build_pptx(path: Path) -> None:
    """Write SLIDES as a title-and-content pptx."""
    from pptx import Presentation

    prs = Presentation()
    for title, body in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(str(path))


def answer_key_markdown() -> str:
    """Human-readable answer key for comparison experiments."""
    lines = [
        "# 違反を仕込んだサンプル資料: 正解表",
        "",
        "`violation-deck.pptx`(架空の提案書 9 ページ)に仕込んだ 20 件の違反。",
        "基準パック: `backend/tests/fixtures/packs/sample_care_sales.yaml`",
        "",
        "| # | ルール | ページ(0=資料全体) |",
        "|---|---|---|",
    ]
    lines += [f"| {i} | {rule} | {page} |" for i, (rule, page) in enumerate(EXPECTED, start=1)]
    return "\n".join(lines) + "\n"


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    out.mkdir(parents=True, exist_ok=True)
    build_pptx(out / "violation-deck.pptx")
    (out / "answer-key.md").write_text(answer_key_markdown(), encoding="utf-8")
    print(f"wrote {out / 'violation-deck.pptx'} and {out / 'answer-key.md'}")

    # PDF 版(Render は PDF のみ受け付ける)。LibreOffice が無ければ作らない
    from app.infra.converter import default_soffice, pptx_to_pdf, soffice_available
    from app.infra.pptx_fix import fill_theme_east_asian_fonts

    soffice = default_soffice()
    if soffice_available(soffice):
        data = fill_theme_east_asian_fonts((out / "violation-deck.pptx").read_bytes())
        (out / "violation-deck.pdf").write_bytes(pptx_to_pdf(data, soffice))
        print(f"wrote {out / 'violation-deck.pdf'}")
    else:
        print("LibreOffice が無いため PDF 版は作りませんでした")
